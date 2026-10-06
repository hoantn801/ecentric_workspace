# Copyright (c) 2026, eCentric and contributors
"""Tab SLA + Phan bo cong viec tren /tong-quan#nhan-su (Hoan 01/10). Thuan Python, KHONG can bench:
python -m unittest ecentric_workspace.hr.tests.test_overview_team_summary"""
import datetime
import os
import sys
import unittest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..")))

from ecentric_workspace.hr.overview import team_summary as TS  # noqa: E402

NOW = datetime.datetime(2026, 10, 1, 9, 0)
PAST = datetime.datetime(2026, 9, 20, 18, 0)
FUTURE = datetime.datetime(2026, 10, 30, 18, 0)
DEPTS = {"Sales - EC": {"label": "Sales", "manager": "Trưởng Sales"}, "Ops - EC": {"label": "Ops", "manager": None}}


def ob(user, status, dept="Sales - EC", group="attendance", due=PAST, counts=1):
    return {"owner_user": user, "department": dept, "group_key": group, "counts_toward_sla": counts,
            "status": status, "due_at": due, "paused_seconds": 0}


class TestSla(unittest.TestCase):
    def build(self, rows, people=None):
        people = people if people is not None else {
            "a@x": {"name": "An", "department": "Sales - EC", "active": True},
            "b@x": {"name": "Binh", "department": "Sales - EC", "active": True},
            "c@x": {"name": "Chi", "department": "Ops - EC", "active": True}}
        return TS.build_sla(rows, people, DEPTS, "2026-09", NOW)

    def test_ti_le_nguoi_phong_cong_ty(self):
        rows = [ob("a@x", "Met")] * 9 + [ob("a@x", "Late")] + [ob("b@x", "Met"), ob("b@x", "Missed")]
        d = self.build(rows)
        sales = [x for x in d["departments"] if x["department"] == "Sales - EC"][0]
        a = [m for m in sales["members"] if m["user"] == "a@x"][0]
        b = [m for m in sales["members"] if m["user"] == "b@x"][0]
        self.assertEqual((a["rate"], a["band"]), (90.0, "good"))
        self.assertEqual((b["rate"], b["band"]), (50.0, "bad"))
        # phong = tong dung han / tong da cham, KHONG phai trung binh cua ti le tung nguoi
        self.assertEqual(sales["rate"], round(10 * 100.0 / 12, 1))
        self.assertEqual(sales["manager"], "Trưởng Sales")
        self.assertEqual(d["stats"]["rate"], round(10 * 100.0 / 12, 1))
        self.assertEqual((d["stats"]["good"], d["stats"]["bad"], d["stats"]["none"]), (1, 1, 1))
        # thap len dau trong phong
        self.assertEqual(sales["members"][0]["user"], "b@x")

    def test_nguoi_chua_co_dau_viec_van_hien_va_xuong_cuoi(self):
        d = self.build([ob("a@x", "Met")])
        ops = [x for x in d["departments"] if x["department"] == "Ops - EC"][0]
        self.assertEqual(ops["members"][0]["rate"], None)
        self.assertEqual(ops["members"][0]["band"], "none")
        self.assertEqual(d["departments"][-1]["department"], "Ops - EC")

    def test_viec_con_han_khong_tinh(self):
        d = self.build([ob("a@x", "Met"), ob("a@x", "Open", due=FUTURE)], people={})
        m = d["departments"][0]["members"][0]
        self.assertEqual((m["rate"], m["scored"], m["open"]), (100.0, 1, 1))

    def test_nhom_ngoai_sla_khong_vao_tong(self):
        # approval thang 9 ngoai %SLA (GROUP_OFF_PERIODS) - giong /sla: dong mang counts=0
        d = self.build([ob("a@x", "Met"), ob("a@x", "Late", group="approval", counts=0)])
        m = [x for x in d["departments"] if x["department"] == "Sales - EC"][0]["members"]
        a = [x for x in m if x["user"] == "a@x"][0]
        self.assertEqual(a["rate"], 100.0)
        appr = [g for g in d["groups"] if g["key"] == "approval"][0]
        self.assertFalse(appr["counts"])
        self.assertEqual(appr["rate"], 0.0)

    def test_phong_theo_nghia_vu_va_nguoi_da_nghi(self):
        d = self.build([ob("z@x", "Met", dept="Ops - EC")], people={})
        self.assertEqual(d["departments"][0]["department"], "Ops - EC")
        m = d["departments"][0]["members"][0]
        self.assertTrue(m["left"])
        self.assertEqual(m["name"], "z")

    def test_khong_co_truong_luong(self):
        d = self.build([ob("a@x", "Met")])
        self.assertNotIn("salary", repr(d).lower())
        self.assertNotIn("luong", repr(d).lower())


def doc(name, emp, status=None, level=None, creation="2026-10-01 08:00", req=None, dept="Sales - EC"):
    return {"name": name, "employee": emp, "employee_name": emp.upper(), "department": dept, "creation": creation,
            "approval_request": req or ("REQ-" + name), "approval_status": status, "current_level": level}


class TestBrand(unittest.TestCase):
    LABELS = {"FLD-VN": "France Lait", "ABC": "Abc", "NP": "New Project"}

    def expected(self):
        return [{"employee": "E1", "name": "An", "department": "Sales - EC"},
                {"employee": "E2", "name": "Binh", "department": "Sales - EC"},
                {"employee": "E3", "name": "Chi", "department": "Ops - EC"},
                {"employee": "E4", "name": "Dung", "department": "Ops - EC"}]

    def test_trang_thai_va_ty_trong(self):
        docs = [doc("D1", "E1", "Approved"), doc("D2", "E2", "Pending", 1),
                doc("D3", "E3", "Information Required", 1, dept="Ops - EC"),
                doc("D0", "E1", "Rejected", creation="2026-10-01 09:00")]
        details = {"D1": {"FLD-VN": 60.0, "ABC": 40.0}, "D2": {"ABC": 100.0}, "D3": {"NP": 100.0},
                   "D0": {"NP": 100.0}}
        d = TS.build_brand(self.expected(), docs, details, {"REQ-D2": ["Lead A"]}, self.LABELS, DEPTS, "2026-09")
        st = d["stats"]
        self.assertEqual((st["total"], st["final"], st["wait_lead"], st["returned"], st["none"]), (4, 1, 1, 1, 1))
        # phieu Rejected moi hon KHONG che phieu Approved song
        sales = [x for x in d["departments"] if x["department"] == "Sales - EC"][0]
        e1 = [m for m in sales["members"] if m["employee"] == "E1"][0]
        self.assertEqual((e1["status"], e1["doc"]), ("final", "D1"))
        e2 = [m for m in sales["members"] if m["employee"] == "E2"][0]
        self.assertEqual(e2["waiting_on"], ["Lead A"])
        # ty trong chung chi tu phieu da chot
        self.assertEqual([(b["label"], b["fte"], b["share"]) for b in d["mix"]],
                         [("France Lait", 0.6, 60.0), ("Abc", 0.4, 40.0)])
        self.assertEqual(sales["state"], "partial")
        ops = [x for x in d["departments"] if x["department"] == "Ops - EC"][0]
        self.assertEqual((ops["state"], ops["returned"], ops["none"]), ("partial", 1, 1))
        self.assertEqual(ops["mix"], [])
        # phong chua xong len dau
        self.assertNotEqual(d["departments"][-1]["state"], "none")
        self.assertEqual({b["id"] for b in d["brands"]}, {"FLD-VN", "ABC", "NP"})

    def test_cho_truong_phong(self):
        d = TS.build_brand(self.expected()[:1], [doc("D1", "E1", "Pending", 2)], {"D1": {"ABC": 100.0}},
                           {"REQ-D1": ["TP"]}, self.LABELS, DEPTS, "2026-09")
        self.assertEqual(d["stats"]["wait_head"], 1)
        self.assertEqual(d["departments"][0]["members"][0]["waiting_on"], ["TP"])

    def test_nguoi_da_nghi_co_phieu_van_hien(self):
        d = TS.build_brand([], [doc("D9", "E9", "Approved")], {"D9": {"ABC": 100.0}}, {}, self.LABELS, DEPTS, "2026-09")
        self.assertEqual(d["stats"]["final"], 1)
        self.assertEqual(d["departments"][0]["members"][0]["name"], "E9")

    def test_ca_phong_chot_xong(self):
        exp = self.expected()[:2]
        docs = [doc("D1", "E1", "Approved"), doc("D2", "E2", "Approved")]
        d = TS.build_brand(exp, docs, {"D1": {"ABC": 50.0, "NP": 50.0}, "D2": {"ABC": 100.0}}, {}, self.LABELS, DEPTS, "2026-09")
        s = d["departments"][0]
        self.assertEqual(s["state"], "done")
        self.assertEqual([(b["id"], b["fte"], b["share"]) for b in s["mix"]], [("ABC", 1.5, 75.0), ("NP", 0.5, 25.0)])


class TestXuatExcel(unittest.TestCase):
    def test_sla_ba_sheet(self):
        d = TS.build_sla([ob("a@x", "Met"), ob("a@x", "Late")], {"a@x": {"name": "An", "department": "Sales - EC", "active": True}},
                         DEPTS, "2026-09", NOW)
        name, sheets = TS.sla_sheets(d)
        self.assertEqual(name, "SLA_thang_9-2026.xlsx")
        self.assertEqual([t for t, _ in sheets], ["Theo người", "Theo phòng", "Theo nhóm việc"])
        people = sheets[0][1]
        self.assertEqual(people[1][:6], ["Sales", "An", "a@x", 50.0, 1, 2])
        self.assertEqual(len(people[0]), len(people[1]))
        self.assertEqual(sheets[1][1][-1][0], "Toàn công ty")

    def test_brand_ba_sheet_va_tong(self):
        exp = [{"employee": "E1", "name": "An", "department": "Sales - EC"},
               {"employee": "E2", "name": "Binh", "department": "Sales - EC"}]
        d = TS.build_brand(exp, [doc("D1", "E1", "Approved")], {"D1": {"ABC": 70.0, "NP": 30.0}}, {},
                           {"ABC": "Abc", "NP": "New Project"}, DEPTS, "2026-09")
        name, sheets = TS.brand_sheets(d)
        self.assertEqual(name, "Phan_bo_cong_viec_thang_9-2026.xlsx")
        people = sheets[0][1]
        self.assertEqual(people[0], ["Phòng", "Nhân viên", "Mã NV", "Trạng thái", "Đang chờ", "Abc (%)", "New Project (%)", "Tổng (%)"])
        self.assertEqual(people[1], ["Sales", "An", "E1", "Đã chốt", "", 70.0, 30.0, 100.0])
        self.assertEqual(people[2][3], "Chưa nộp")
        self.assertEqual(sheets[1][1][-1], ["Tổng (phiếu đã chốt)", 1.0, 100])

    def test_ghi_xlsx_doc_lai_duoc(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("khong co openpyxl")
        import importlib.util
        import io
        spec = importlib.util.spec_from_file_location(
            "ec_xlsx_export", os.path.join(os.path.dirname(__file__), "..", "overview", "xlsx_export.py"))
        X = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(X)
        data = X.build([("Theo người", [["Phòng", "%"], ["Sales", 92.5], ["Ops", None]]), ("Trống", [])])
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ["Theo người", "Trống"])
        ws = wb["Theo người"]
        self.assertEqual((ws["A1"].value, ws["B2"].value, ws["B3"].value), ("Phòng", 92.5, None))
        self.assertTrue(ws["A1"].font.bold)
        self.assertEqual(ws.freeze_panes, "A2")


class TestTrang(unittest.TestCase):
    def test_trang_co_hai_tab_trong_script_san_co(self):
        p = os.path.join(os.path.dirname(__file__), "..", "pages", "tong_quan", "main_section.html")
        with open(p, encoding="utf-8") as fh:
            h = fh.read()
        import re
        self.assertEqual(re.findall(r'<script[^>]*\bid="([^"]+)"', h), ["ec-tongquan-js"])
        self.assertIn('["sla", "SLA"], ["phan-bo", "Phân bổ công việc"]', h)
        self.assertIn("ec-tq-sla-brand-v1", h)
        self.assertIn("get_sla_summary", h)
        self.assertIn("get_brand_summary", h)
        self.assertIn("export_summary_xlsx", h)
        self.assertNotIn("csvDownload", h)


if __name__ == "__main__":
    unittest.main()
