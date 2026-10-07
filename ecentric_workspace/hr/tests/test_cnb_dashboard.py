# Copyright (c) 2026, eCentric and contributors
"""ec-cnb-dashboard-v1 (07/10/2026): file "Timesheet Thang N Dashboard" cua CnB xuat thang tu ERP.

Ham thuan (chi openpyxl) - chay khong can bench:
    python -m unittest ecentric_workspace.hr.tests.test_cnb_dashboard
Da doi chieu tay voi file CnB thang 9/2026: 78 dong bang cong, SLA va % Brand trung tung o;
dashboard trung so tru 2 cho file tay sai (tong phep chi cong 30 dong dau, ty le chia doi)."""
import importlib.util
import io
import os
import unittest

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
_spec = importlib.util.spec_from_file_location(
    "_cnb_dashboard", os.path.join(HERE, "..", "overview", "cnb_dashboard.py"))
CD = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(CD)


def _report(n_emp=3, ndays=30):
    cols = [{"fieldname": "employee"}, {"fieldname": "employee_name"}, {"fieldname": "ec_nv_chot_cong"},
            {"fieldname": "ec_lead_chot_cong"}, {"fieldname": "shift"}]
    cols += [{"fieldname": "%02d-09-2026" % (d + 1), "label": "%d X" % (d + 1)} for d in range(ndays)]
    data = []
    for i in range(n_emp):
        row = {"employee": "E%d" % i, "employee_name": "Nguoi %d" % i, "ec_nv_chot_cong": "Đã chốt 01/10 09:00",
               "ec_lead_chot_cong": "Chưa chốt", "shift": "EC Standard 9-18"}
        for d in range(ndays):
            row["%02d-09-2026" % (d + 1)] = ["P", "WO", "L", "HD/P", "H", "A", ""][d % 7]
        data.append(row)
    return cols, data, None, None


SLA = {"groups": [{"key": "g1", "label": "Chấm công"}],
       "departments": [{"label": "Service", "members": [
           {"user": "b@x", "name": "B", "rate": 95.0, "ontime": 19, "scored": 20, "late": 1, "missed": 0, "open": 0,
            "left": False, "groups": {"g1": {"rate": 95.0}}},
           {"user": "a@x", "name": "A", "rate": None, "ontime": 0, "scored": 0, "late": 0, "missed": 0, "open": 0,
            "left": True, "groups": {}}]}]}
BW = {"brands": [{"id": "B1", "label": "Brand Một"}],
      "departments": [{"label": "Service", "members": [
          {"employee": "E1", "name": "Nguoi 1", "status": "final", "weights": {"B1": 100.0}, "waiting_on": []}]}]}


def _build(n_emp=3, depts=("Service", "Production", "Service")):
    days, ts = CD.timesheet_rows(_report(n_emp), {"E%d" % i: depts[i % len(depts)] for i in range(n_emp)})
    sg, sla = CD.sla_rows(SLA, {"a@x": "E0", "b@x": "E1"})
    bb, bw = CD.brand_rows(BW)
    return openpyxl.load_workbook(io.BytesIO(CD.build("2026-09", days, ts, sg, sla, bb, bw)))


class CnbDashboard(unittest.TestCase):
    def test_dung_4_sheet_va_ten_nhu_file_tay(self):
        wb = _build()
        self.assertEqual(wb.sheetnames, ["Tổng Quan", "Timesheet T09.2026", "SLA T09.2026", "% Brand"])
        self.assertEqual(CD.filename("2026-09"), "Timesheet_Thang_9_2026_Dashboard.xlsx")

    def test_bang_cong_giu_nguyen_ma_va_cong_thuc_tong(self):
        ws = _build()["Timesheet T09.2026"]
        head = [ws.cell(1, c).value for c in range(1, 42)]
        self.assertEqual(head[:6], CD.TS_FIXED)
        self.assertEqual(head[6], "1 X")
        self.assertEqual(head[36:], CD.TS_TOTALS)
        self.assertEqual([ws.cell(2, c).value for c in range(2, 8)],
                         ["E0", "Nguoi 0", "Đã chốt 01/10 09:00", "Chưa chốt", "EC Standard 9-18", "P"])
        self.assertIsNone(ws["M2"].value)   # ngay khong co ban ghi -> o trong, khong ghi chuoi rong
        self.assertEqual(ws["AK2"].value, '=COUNTIF(G2:AJ2,"P")+COUNTIF(G2:AJ2,"HD/P")*0.5')
        self.assertEqual(ws["AN3"].value, '=COUNTIF(G3:AJ3,"A")+COUNTIF(G3:AJ3,"HD/A")*0.5')
        self.assertEqual(ws["AO4"].value, "=AK4+AL4")
        self.assertEqual(ws.auto_filter.ref, "A1:AO4")

    def test_thang_28_ngay_doi_cot_tong(self):
        days, ts = CD.timesheet_rows(_report(1, 28), {})
        wb = openpyxl.load_workbook(io.BytesIO(CD.build("2026-02", days, ts, [], [], [], [])))
        ws = wb["Timesheet T02.2026"]
        self.assertEqual(ws["AI1"].value, "Công thực tế (P)")
        self.assertIn("SUM('Timesheet T02.2026'!AI2:AI2)", wb["Tổng Quan"]["C8"].value)

    def test_dashboard_tro_toi_dong_cuoi_khong_cat_30_dong(self):
        ov = _build(n_emp=40)["Tổng Quan"]
        for cell in ("B8", "C8", "D8", "E8", "F8", "D13", "D14", "D15", "D16"):
            f = ov[cell].value
            self.assertIn("41", f, cell)
            self.assertNotIn("31", f, cell)
        self.assertEqual(ov["E8"].value,
                         "=IF(C5=\"All Department\", SUM('Timesheet T09.2026'!AM2:AM41), "
                         "SUMIF('Timesheet T09.2026'!A2:A41, C5, 'Timesheet T09.2026'!AM2:AM41))")
        self.assertEqual(ov["D17"].value, "=D13+D14+D15")
        self.assertEqual(ov["E13"].value, "=IF(SUM(D$13:D$16)=0,0,D13/SUM(D$13:D$16))")

    def test_o_chon_phong_ban_lay_tu_du_lieu(self):
        ov = _build()["Tổng Quan"]
        self.assertEqual(ov["C5"].value, "All Department")
        dv = ov.data_validations.dataValidation
        self.assertEqual(len(dv), 1)
        self.assertEqual(dv[0].formula1, '"All Department,Production,Service"')
        self.assertIn("C5", str(dv[0].sqref))

    def test_sla_va_brand_them_phong_ban_ma_nv(self):
        wb = _build()
        s = wb["SLA T09.2026"]
        self.assertEqual([s.cell(1, c).value for c in range(1, 13)], CD.SLA_HEAD + ["% Chấm công"])
        self.assertEqual([s.cell(2, c).value for c in range(1, 5)], ["Service", "E0", "A", "a@x"])
        self.assertEqual(s["K2"].value, "x")
        self.assertEqual(s.freeze_panes, "A2")
        b = wb["% Brand"]
        self.assertEqual([b.cell(1, c).value for c in range(1, 9)], CD.BW_HEAD + ["Brand Một (%)", "Tổng (%)"])
        self.assertEqual([b.cell(2, c).value for c in range(1, 9)],
                         ["Service", "E1", "Nguoi 1", "E1", "Đã chốt", None, 100, 100])

    def test_mau_va_dinh_dang_nhu_file_tay(self):
        wb = _build()
        ws = wb["Timesheet T09.2026"]
        self.assertEqual(ws["A1"].fill.fgColor.rgb, "FF1E293B")
        self.assertTrue(ws["A1"].font.b)
        self.assertEqual(ws["A2"].font.name, "Arial")
        self.assertEqual(ws["A3"].fill.fgColor.rgb, "FFF8FAFC")   # soc dong le
        self.assertIsNone(ws["A2"].fill.fill_type)
        rules = [str(cf.sqref) for cf in ws.conditional_formatting]
        self.assertEqual(rules, ["D2:D4", "E2:E4", "G2:AO4"])
        ov = wb["Tổng Quan"]
        self.assertIn("B2:F3", [str(r) for r in ov.merged_cells.ranges])
        self.assertEqual(ov["B2"].value, "BÁO CÁO TỔNG QUAN CHẤM CÔNG THÁNG 9/2026")


if __name__ == "__main__":
    unittest.main()
