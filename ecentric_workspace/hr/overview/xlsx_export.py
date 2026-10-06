# Copyright (c) 2026, eCentric and contributors
"""Ghi nhieu sheet ra .xlsx (openpyxl co san trong Frappe). Thuan: nhan [(ten sheet, hang)],
tra ve bytes. Hang dau moi sheet la tieu de: in dam, co dinh khi cuon, co bo loc."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_HEAD_FILL = PatternFill("solid", fgColor="EEF0FB")


def build(sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title=str(title)[:31])
        for r in rows:
            ws.append(["" if c is None else c for c in r])
        if not rows:
            continue
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = _HEAD_FILL
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i in range(1, len(rows[0]) + 1):
            width = max(len(str(r[i - 1])) if i - 1 < len(r) and r[i - 1] is not None else 0 for r in rows)
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 45)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
